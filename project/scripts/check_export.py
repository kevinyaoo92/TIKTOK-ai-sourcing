# -*- coding: utf-8 -*-
"""Read newest exported xlsx under data/raw/tiktok/thailand and print a structure summary."""
import glob, os, sys, json
from openpyxl import load_workbook

SAVE_DIR = r"D:\tiktok-ai-sourcing\project\data\raw\tiktok\thailand"
files = sorted(glob.glob(os.path.join(SAVE_DIR, "*.xlsx")), key=os.path.getmtime, reverse=True)
if not files:
    print(json.dumps({"ok": False, "error": "no xlsx found in " + SAVE_DIR}, ensure_ascii=False))
    sys.exit(1)

f = files[0]
wb = load_workbook(f, data_only=True, read_only=False)
ws = wb.active
rows = []
for row in ws.iter_rows(values_only=True):
    rows.append([("" if c is None else str(c).replace("\n", "").strip()) for c in row])
wb.close()
while rows and all(not c for c in rows[-1]):
    rows.pop()

meta = {}
hdr_idx = None
for i, r in enumerate(rows):
    if hdr_idx is not None:
        break
    nonempty = [c for c in r if c]
    if nonempty and nonempty[0].startswith("["):
        key, _, val = nonempty[0][1:].partition("]:")
        meta[key.strip()] = val.strip()
        continue
    hdr_idx = i  # first non-meta row (blank or header)
# header = first row whose col0 == 关键词; otherwise rows[hdr_idx]
hdr_idx2 = next((i for i, r in enumerate(rows) if r and r[0] == "关键词"), hdr_idx)
header = rows[hdr_idx2]
data = [r for r in rows[hdr_idx2 + 1:] if any(c for c in r)]

summary = {
    "ok": True,
    "file_path": f,
    "file_format": os.path.splitext(f)[1].lstrip("."),
    "sheet": ws.title,
    "total_rows": len(rows),
    "data_rows": len(data),
    "cols": len(header),
    "columns": header,
    "first_row": data[0] if data else [],
    "last_row": data[-1] if data else [],
    "meta": meta,
}
print(json.dumps(summary, ensure_ascii=False, indent=1))